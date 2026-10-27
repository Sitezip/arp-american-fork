#!/usr/bin/env python3
"""Convert the AF ARP meeting schedule spreadsheet into data/meetings.json.

Usage:
    pip install openpyxl
    python3 scripts/xlsx_to_json.py "Meeting Schedule - AF.xlsx" [--sheet Regular] [--out data/meetings.json]

What it does
- Reads ONE sheet (default "Regular", else the first sheet). It never reads the
  "With Facilitators" sheet: that sheet contains participant first names that
  must NOT be published.
- Finds each section by its header row (column B == "Audience"), e.g.
  "General Addiction - In-person Meetings" or "General Addiction - ZOOM Meetings".
- Every row under a section whose Day column is a weekday becomes a meeting.
- Values are copied as written (whitespace tidied only). Times are kept as
  written in "time"; a parsed 24h "time_sort" is added only for sorting.
- The schedule date is taken from the first date-valued cell in the sheet
  (the spreadsheet keeps it in the top-right of the first header row).
- The Zoom dial-in number is read from the sheet text ("call 253-...").

Review the diff of data/meetings.json before committing.
"""
import argparse, datetime as dt, json, re, sys, urllib.parse
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Please install openpyxl: pip install openpyxl")

DAYS = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
CITIES = ["American Fork", "Pleasant Grove", "Cedar Hills", "Highland", "Lehi", "Alpine", "Saratoga Springs"]


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, dt.time):
        return v.strftime("%-I:%M")
    return re.sub(r"\s+", " ", str(v)).strip()


def time_sort(raw):
    """Return 'HH:MM' (24h) for sorting, or '' if it can't be parsed. Display uses the raw value."""
    t = raw.lower().replace(".", "").strip()
    if t == "noon":
        return "12:00"
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?\s*(am|pm)?$", t)
    if not m:
        return ""
    h, mnt, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if ap == "pm" and h != 12:
        h += 12
    elif ap == "am" and h == 12:
        h = 0
    elif ap is None and 1 <= h <= 11:
        h += 12  # ARP evening meetings written without am/pm (e.g. "7:30") are assumed PM for sorting only
    return f"{h:02d}:{mnt:02d}"


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="Regular")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent.parent / "data" / "meetings.json"))
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    ws = wb[a.sheet] if a.sheet in wb.sheetnames else wb.worksheets[0]
    if "facilitator" in ws.title.lower():
        sys.exit("Refusing to read a 'With Facilitators' sheet (contains participant names).")

    rows = [[c.value for c in r] for r in ws.iter_rows(max_col=10)]
    updated, phone = None, ""
    for r in rows:
        for v in r:
            if updated is None and isinstance(v, dt.datetime):
                updated = v.date().isoformat()
            if isinstance(v, str) and not phone:
                m = re.search(r"call\s+([\d\-\(\) ]{10,})", v, re.I)
                if m:
                    phone = m.group(1).strip()

    meetings, section, is_zoom, seen = [], "", False, {}
    for r in rows:
        r = (r + [None] * 10)[:10]
        a_, b, c, d, e, f, g, h = (clean(x) for x in r[:8])
        if b == "Audience":
            section = a_
            is_zoom = "zoom" in a_.lower() or e.lower().startswith("zoom")
            continue
        if not a_ or c not in DAYS:
            continue
        type_ = section.split(" - ")[0].strip() if section else ""
        city = next((x for x in CITIES if a_.lower().startswith(x.lower())), "")
        focus = a_.split(" - ", 1)[1] if " - " in a_ else a_
        if is_zoom:
            focus = re.sub(r"^ZOOM/Phone-In\s*", "", focus, flags=re.I)
        m = {
            "name": a_,
            "focus": focus.strip(),
            "type": type_,
            "section": section,
            "format": "zoom" if is_zoom else "in-person",
            "audience": b,
            "day": c,
            "time": d,
            "time_sort": time_sort(d),
            "city": city,
        }
        if is_zoom:
            m.update({"zoom_url": e, "meeting_id": g, "passcode": h, "phone": phone, "location": ""})
        else:
            q = e if (any(x.lower() in e.lower() for x in CITIES) or not city) else f"{e}, {city}"
            if "ut" not in q.lower().split(", ")[-1].split():
                q += ", UT"
            m.update({"location": e, "map_url": "https://www.google.com/maps/search/?api=1&query=" + urllib.parse.quote(q)})
        m["group_leaders"] = f
        m["notes"] = ""
        base = slug(f"{c}-{m['time_sort'] or d}-{a_}-{b}")
        seen[base] = seen.get(base, 0) + 1
        m["id"] = base if seen[base] == 1 else f"{base}-{seen[base]}"
        meetings.append(m)

    out = {
        "_comment": "Generated by scripts/xlsx_to_json.py from the meeting schedule spreadsheet. "
                    "Edit the spreadsheet and re-run, or edit carefully by hand (see CONTENT-EDITING.md).",
        "source_file": Path(a.xlsx).name,
        "source_sheet": ws.title,
        "schedule_updated": updated or dt.date.today().isoformat(),
        "zoom_phone": phone,
        "count": len(meetings),
        "meetings": meetings,
    }
    Path(a.out).write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n")
    print(f"Wrote {len(meetings)} meetings to {a.out} (schedule date {out['schedule_updated']})")


if __name__ == "__main__":
    main()
